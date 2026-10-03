"""
tests/test_fastapi_server.py
----------------------------
Unit and integration tests for InsiEDR FastAPI ASGI Full-Stack REST + SSE API:
- Full-Stack REST v1 endpoints and dual-route aliases
- Server-Sent Events (SSE) streaming engine
- Streaming chunked export (CSV / JSON)
- Interactive Swagger UI & OpenAPI specification
- Centralized RFC 7807 JSON error handling
- In-memory TTL caching
- Agent telemetry encryption & ingestion
"""
from __future__ import annotations

import json
import pytest
from unittest.mock import MagicMock
from fastapi.testclient import TestClient

from server.app import create_app
from server.api.cache import api_cache
from server.api.events import broadcaster, dispatch_threat_event, dispatch_agent_event


class MockStorage:
    """Mock storage adapter for API endpoint unit tests."""

    def __init__(self):
        self.stats_calls = 0

    def get_stats(self):
        self.stats_calls += 1
        return {
            "agents": 4,
            "logs": 1200,
            "anomalies": 15,
            "risk_events": 42,
            "baselines": 4,
        }

    def list_agents(self, limit=100, offset=0):
        return [
            {"agent_id": "agent-01", "hostname": "DESKTOP-SEC01", "status": "active", "os_system": "Windows"},
            {"agent_id": "agent-02", "hostname": "DESKTOP-DEV02", "status": "active", "os_system": "Windows"},
        ]

    def list_logs(self, limit=100, offset=0, **kwargs):
        return [
            {
                "id": 1,
                "collected_at": "2026-09-08T10:00:00Z",
                "agent_id": "agent-01",
                "hostname": "DESKTOP-SEC01",
                "username": "alice",
                "collector": "file-collector",
                "status": "success",
                "payload": {"files_accessed": 12},
            },
            {
                "id": 2,
                "collected_at": "2026-09-08T10:01:00Z",
                "agent_id": "agent-02",
                "hostname": "DESKTOP-DEV02",
                "username": "bob",
                "collector": "device-collector",
                "status": "success",
                "payload": {"usb_inserted": True},
            },
        ]

    def list_risk_events(self, limit=100, offset=0):
        return [
            {
                "id": 101,
                "created_at": "2026-09-08T10:05:00Z",
                "payload_id": "payload-abc",
                "agent_id": "agent-01",
                "username": "alice",
                "risk_level": "high",
                "risk_score": 78.5,
                "summary": "Bulk File Access Anomaly",
                "correlated_signals_json": {"heuristics": {"overall_score": 78.5}},
            }
        ]

    def list_anomalies(self, limit=100, offset=0):
        return [
            {
                "id": 1,
                "anomaly_type": "file_bulk_access",
                "severity": "HIGH",
                "username": "alice",
            }
        ]

    def get_pc_status(self, seconds_since_online=300):
        return {"online_count": 2, "offline_count": 0, "total_count": 2}

    def list_normalized_features(self, limit=100, offset=0, **kwargs):
        return [
            {
                "username": "alice",
                "agent_id": "agent-01",
                "collector": "keystroke-collector",
                "feature_name": "mean_flight_time_ms",
                "feature_value": 110.5,
                "window_end": "2026-09-08T10:00:00Z",
                "features_json": {"mean_flight_time_ms": 110.5, "typing_speed_cpm": 320.0},
            },
            {
                "username": "alice",
                "agent_id": "agent-01",
                "collector": "file-collector",
                "feature_name": "files_accessed",
                "feature_value": 15.0,
                "window_end": "2026-09-08T10:00:00Z",
                "features_json": {"files_accessed": 15.0},
            },
        ]

    def list_collector_results(self, limit=100, offset=0, collector=None, username=None, **kwargs):
        return [
            {
                "payload_id": "payload-ks-01",
                "agent_id": "agent-01",
                "hostname": "DESKTOP-SEC01",
                "collector": "keystroke-collector",
                "collector_collected_at": "2026-09-08T10:00:00Z",
                "received_at": "2026-09-08T10:00:01Z",
                "username": "alice",
                "status": "success",
                "payload_json": {
                    "mean_flight_time_ms": 120.5,
                    "std_flight_time_ms": 15.2,
                    "mean_dwell_time_ms": 65.4,
                    "std_dwell_time_ms": 8.1,
                    "typing_speed_cpm": 280.0,
                    "backspace_ratio": 0.04,
                    "keystroke_timings": [[65.4, 120.5], [70.1, 115.2]],
                },
                "features_json": {
                    "mean_flight_time_ms": 120.5,
                    "std_flight_time_ms": 15.2,
                    "mean_dwell_time_ms": 65.4,
                    "std_dwell_time_ms": 8.1,
                    "typing_speed_cpm": 280.0,
                    "backspace_ratio": 0.04,
                },
            }
        ]

    def get_distinct_collectors(self):
        return ["keystroke-collector", "file", "logon", "process"]

    def get_distinct_usernames(self):
        return ["alice", "bob"]


@pytest.fixture
def mock_storage():
    return MockStorage()


@pytest.fixture
def client(mock_storage):
    app = create_app(storage=mock_storage, apply_migrations=False)
    api_cache.clear()
    with TestClient(app, raise_server_exceptions=False) as c:
        yield c


# ------------------------------------------------------------------------------
# 1. OpenAPI & Swagger UI Documentation Tests
# ------------------------------------------------------------------------------
def test_swagger_ui_renders(client):
    """Verify that Swagger UI HTML loads on /docs and /openapi.json."""
    resp = client.get("/docs")
    assert resp.status_code == 200
    assert "swagger-ui" in resp.text.lower()

    resp_json = client.get("/openapi.json")
    assert resp_json.status_code == 200
    spec = resp_json.json()
    assert spec["openapi"].startswith("3.")
    assert "paths" in spec
    assert "/api/logs" in spec["paths"]
    assert "/api/v1/stream/threats" in spec["paths"]
    assert "/api/v1/export/logs" in spec["paths"]


# ------------------------------------------------------------------------------
# 2. Centralized RFC 7807 Error Handling Tests
# ------------------------------------------------------------------------------
def test_rfc7807_not_found_error(client):
    """Verify 404 returns structured JSON envelope, not HTML."""
    resp = client.get("/api/nonexistent-endpoint-test")
    assert resp.status_code == 404
    data = resp.json()
    assert data["ok"] is False
    assert data["success"] is False
    assert "error" in data
    assert data["error"]["code"] == "NOT_FOUND"


# ------------------------------------------------------------------------------
# 3. Dual-Route Compatibility Tests
# ------------------------------------------------------------------------------
def test_dual_route_agents(client):
    """Verify /api/agents and /api/v1/agents return identical data."""
    r1 = client.get("/api/agents")
    r2 = client.get("/api/v1/agents")
    assert r1.status_code == 200
    assert r2.status_code == 200
    assert r1.json() == r2.json()
    assert len(r1.json()["agents"]) == 2


def test_dashboard_summary_and_stats(client):
    """Verify /api/dashboard-summary and /api/stats work for fleet telemetry."""
    r_summary = client.get("/api/dashboard-summary")
    r_v1 = client.get("/api/v1/dashboard-summary")
    assert r_summary.status_code == 200
    assert r_v1.status_code == 200
    data = r_summary.json()
    assert data["ok"] is True
    assert len(data["agents"]) == 2
    assert "risk_counts" in data


# ------------------------------------------------------------------------------
# 4. In-Memory TTL Cache Tests
# ------------------------------------------------------------------------------
def test_stats_ttl_caching(client, mock_storage):
    """Verify /api/stats serves cached results on consecutive calls."""
    assert mock_storage.stats_calls == 0

    # First call: cache miss, hits storage
    r1 = client.get("/api/v1/stats")
    assert r1.status_code == 200
    assert mock_storage.stats_calls == 1

    # Second call: cache hit, does not hit storage
    r2 = client.get("/api/stats")
    assert r2.status_code == 200
    assert mock_storage.stats_calls == 1  # call count did not increment!
    assert r1.json() == r2.json()


# ------------------------------------------------------------------------------
# 5. Streaming Export (CSV / JSON) Tests
# ------------------------------------------------------------------------------
def test_streaming_export_logs_csv(client):
    """Verify chunked streaming CSV export for telemetry logs."""
    resp = client.get("/api/v1/export/logs?format=csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers.get("Content-Type", "")
    assert "attachment; filename=" in resp.headers.get("Content-Disposition", "")
    content = resp.text
    assert "id,collected_at,agent_id,hostname,username,collector,status,payload_json" in content
    assert "DESKTOP-SEC01" in content
    assert "alice" in content


def test_streaming_export_threats_csv(client):
    """Verify chunked streaming CSV export for risk events."""
    resp = client.get("/api/v1/export/threats?format=csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers.get("Content-Type", "")
    assert "attachment; filename=" in resp.headers.get("Content-Disposition", "")
    content = resp.text
    assert "id,created_at,payload_id,agent_id,username,risk_level,risk_score" in content
    assert "HIGH" in content
    assert "Bulk File Access Anomaly" in content


def test_streaming_export_logs_ndjson(client):
    """Verify streaming NDJSON export."""
    resp = client.get("/api/v1/export/logs?format=json")
    assert resp.status_code == 200
    assert "application/x-ndjson" in resp.headers.get("Content-Type", "")
    lines = [line.strip() for line in resp.text.split("\n") if line.strip()]
    assert len(lines) == 2
    parsed_first = json.loads(lines[0])
    assert parsed_first["agent_id"] == "agent-01"


# ------------------------------------------------------------------------------
# 6. Server-Sent Events (SSE) Broadcaster Unit Tests
# ------------------------------------------------------------------------------
def test_sse_broadcaster_pubsub():
    """Verify broadcaster correctly pushes events to subscribed queue."""
    q = broadcaster.subscribe("threats")
    try:
        sample_alert = {
            "payload_id": "test-pid",
            "username": "alice",
            "risk_level": "critical",
            "risk_score": 92.0,
        }
        dispatch_threat_event(sample_alert)
        msg = q.get_nowait()
        assert msg["event"] == "threat_alert"
        assert msg["data"]["risk_level"] == "critical"
        assert msg["data"]["risk_score"] == 92.0
    finally:
        broadcaster.unsubscribe("threats", q)


def test_sse_agent_broadcaster():
    """Verify agent status transitions broadcast to subscribers."""
    q = broadcaster.subscribe("agents")
    try:
        agent_event = {"agent_id": "agent-01", "status": "active"}
        dispatch_agent_event(agent_event)
        msg = q.get_nowait()
        assert msg["event"] == "agent_status"
        assert msg["data"]["agent_id"] == "agent-01"
    finally:
        broadcaster.unsubscribe("agents", q)


# ------------------------------------------------------------------------------
# 7. Deep Health Check & Queue Metrics Tests
# ------------------------------------------------------------------------------
def test_deep_health_check(client):
    """Verify /api/health reports deep readiness."""
    resp = client.get("/api/v1/health")
    assert resp.status_code == 200
    data = resp.json()
    assert "database" in data
    assert "queue_backend" in data
    assert "sse_subscribers" in data
    assert "threats" in data["sse_subscribers"]


def test_queue_metrics_endpoint(client):
    """Verify /api/v1/queue/metrics reports queue status."""
    resp = client.get("/api/v1/queue/metrics")
    assert resp.status_code in (200, 503)


# ------------------------------------------------------------------------------
# 8. Agent Telemetry Ingestion Backward Compatibility Tests
# ------------------------------------------------------------------------------
def test_agent_telemetry_ingestion_backward_compatibility(client, mock_storage):
    """Verify both /api/logs and /api/v1/logs accept AES-GCM encrypted telemetry and return 202 Accepted."""
    from datetime import datetime, timezone
    import os
    import uuid
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from shared.crypto_utils import b64encode
    from server.plugin_registry import registry
    from server.crypto.aesgcm_plugin import AESGCMPlugin, AES_GCM_AAD

    # Register test AES-GCM plugin with a random 32-byte key
    test_key = os.urandom(32)
    registry._plugins["aes-256-gcm"] = AESGCMPlugin(test_key)
    aesgcm_cipher = AESGCM(test_key)

    stored_payloads = []
    mock_storage.store_raw_payload = lambda env, payload: stored_payloads.append(payload)

    for endpoint in ("/api/logs", "/api/v1/logs"):
        payload_id = str(uuid.uuid4())
        now_str = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
        telemetry_payload = {
            "schema": "insiedr.agent.telemetry.v1",
            "protocol_version": "2.0",
            "payload_id": payload_id,
            "agent_id": "test-agent-01",
            "collected_at": now_str,
            "hostname": "DESKTOP-TEST",
            "username": "alice",
            "collectors": [
                {
                    "collector": "file-collector",
                    "collected_at": now_str,
                    "hostname": "DESKTOP-TEST",
                    "status": "success",
                    "payload": {"files_modified": 5},
                }
            ],
            "summary": {
                "collector_count": 1,
                "success_count": 1,
                "failed_count": 0,
            },
        }

        # Real AES-256-GCM encryption
        nonce = os.urandom(12)
        plaintext_bytes = json.dumps(telemetry_payload).encode("utf-8")
        ciphertext_bytes = aesgcm_cipher.encrypt(nonce, plaintext_bytes, AES_GCM_AAD)

        envelope = {
            "protocol_version": "2.0",
            "scheme": "aes-256-gcm",
            "payload_id": payload_id,
            "key_id": "k1",
            "nonce": b64encode(nonce),
            "ciphertext": b64encode(ciphertext_bytes),
            "created_at": now_str,
        }

        headers = {
            "X-Crypto-Scheme": "aes-256-gcm",
            "X-Protocol-Version": "2.0",
            "X-Agent-ID": "test-agent-01",
            "X-Payload-ID": payload_id,
            "X-Key-ID": "k1",
            "Content-Type": "application/json",
        }

        resp = client.post(endpoint, content=json.dumps(envelope), headers=headers)
        if resp.status_code != 202:
            print("INGEST ERROR:", resp.status_code, resp.text)
        assert resp.status_code == 202
        body = resp.json()
        assert body["ok"] is True
        assert body["status"] == "accepted"
        assert body["payload_id"] == payload_id

    assert len(stored_payloads) == 2


def test_dashboard_serving(client):
    """Verify that the compiled React SPA dashboard is served correctly at /dashboard/."""
    resp = client.get("/dashboard/")
    assert resp.status_code == 200
    assert "text/html" in resp.headers.get("content-type", "")
    assert "<div id=\"root\"></div>" in resp.text


def test_extension_based_exports(client):
    """Verify direct file extension streaming exports (/api/v1/export/logs.csv and threats.csv)."""
    resp_logs = client.get("/api/v1/export/logs.csv")
    assert resp_logs.status_code == 200
    assert "text/csv" in resp_logs.headers.get("content-type", "")
    assert "id,collected_at,agent_id" in resp_logs.text

    resp_threats = client.get("/api/v1/export/threats.csv")
    assert resp_threats.status_code == 200
    assert "text/csv" in resp_threats.headers.get("content-type", "")
    assert "id,created_at,payload_id" in resp_threats.text


def test_training_dataset_export_csv(client):
    """Verify ML training dataset export in streaming CSV format."""
    resp = client.get("/api/v1/export/training-dataset.csv?days=14&username=alice")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers.get("content-type", "")
    content = resp.text
    assert "user,collected_at,parameters,raw_logs,agent_id,hostname,payload_id" in content
    # Verify user row exists
    assert "alice" in content
    # Verify parameters JSON is present
    assert "{" in content


def test_training_dataset_export_excel(client):
    """Verify ML training dataset export in Excel spreadsheet (.xlsx) format."""
    import io
    import openpyxl

    resp = client.get("/api/v1/export/training-dataset.xlsx?days=30&username=alice")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers.get("content-type", "")
    assert len(resp.content) > 1000

    # Parse in-memory with openpyxl to ensure sheet structure and integrity
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Dataset_Overview" in wb.sheetnames
    assert "Tabular_Feature_Matrix" in wb.sheetnames
    assert "Dataset_Metadata" in wb.sheetnames

    overview_sheet = wb["Dataset_Overview"]
    headers = [cell.value for cell in overview_sheet[1]]
    assert "User" in headers
    assert "Parameters (Model Features)" in headers
    assert "Raw Logs" in headers


def test_training_dataset_export_ndjson(client):
    """Verify ML training dataset export in NDJSON format."""
    resp = client.get("/api/v1/export/training-dataset.json?days=7")
    assert resp.status_code == 200
    assert "application/x-ndjson" in resp.headers.get("content-type", "")
    lines = resp.text.strip().splitlines()
    assert len(lines) > 0
    first_record = json.loads(lines[0])
    assert "user" in first_record
    assert "parameters" in first_record
    assert "raw_logs" in first_record


def test_features_export_csv(client):
    """Verify raw and normalized features export in streaming CSV format."""
    resp = client.get("/api/v1/export/features?format=csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers.get("content-type", "")
    content = resp.text
    assert "id,payload_id,collected_at,agent_id,username,hostname,collector,feature_name,value,source_quality" in content
    assert "alice" in content
    assert "keystroke-collector" in content


def test_features_export_excel(client):
    """Verify normalized features export in Excel spreadsheet (.xlsx) format."""
    import io
    import openpyxl

    resp = client.get("/api/v1/export/features.xlsx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers.get("content-type", "")
    assert len(resp.content) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Normalized_Features" in wb.sheetnames
    sheet = wb["Normalized_Features"]
    headers = [cell.value for cell in sheet[1]]
    assert "Username" in headers
    assert "Collector" in headers
    assert "Feature Name" in headers
    assert "Numeric Value" in headers


def test_keystrokes_export_csv(client):
    """Verify keystroke biometric features export in streaming CSV format."""
    resp = client.get("/api/v1/export/keystrokes.csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers.get("content-type", "")
    content = resp.text
    assert "username,hostname,agent_id,collected_at,mean_flight_time_ms" in content
    assert "120.5" in content
    assert "typing_speed_cpm" in content


def test_keystrokes_export_excel(client):
    """Verify keystroke dynamics export in Excel spreadsheet (.xlsx) format."""
    import io
    import openpyxl

    resp = client.get("/api/v1/export/keystrokes?format=xlsx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers.get("content-type", "")
    assert len(resp.content) > 1000

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "Keystroke_Dynamics" in wb.sheetnames
    kd_sheet = wb["Keystroke_Dynamics"]
    headers = [cell.value for cell in kd_sheet[1]]
    assert "Username" in headers
    assert "Mean Flight (ms)" in headers
    assert "Mean Dwell (ms)" in headers
    assert "Typing Speed (CPM)" in headers

