from __future__ import annotations

import anyio
import csv
import io
import json
import re
from datetime import datetime, timedelta, timezone
from typing import Any, Generator
from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import JSONResponse, Response, StreamingResponse
from server.api.deps import get_storage, require_operator, OperatorPrincipal

router = APIRouter(prefix="/api", tags=["Streaming Export"])
bp = router  # Backward compatibility alias

CHUNK_BATCH_SIZE = 1000


def _format_timestamp(val: Any) -> str:
    if isinstance(val, datetime):
        return val.isoformat()
    return str(val or "")


def _stream_csv_logs(storage, filters: dict[str, Any], limit: int) -> Generator[str, None, None]:
    """Generator yielding CSV rows in chunks for telemetry logs."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # Header
    writer.writerow([
        "id",
        "collected_at",
        "agent_id",
        "hostname",
        "username",
        "collector",
        "status",
        "payload_json",
    ])
    yield buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)

    offset = 0
    exported = 0

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_logs(limit=batch_limit, offset=offset, **filters)
        if not rows:
            break

        for row in rows:
            payload_str = ""
            if row.get("payload"):
                payload_str = json.dumps(row["payload"], default=str)
            elif row.get("payload_json"):
                payload_str = str(row["payload_json"])
            elif row.get("encrypted_envelope_json"):
                payload_str = str(row["encrypted_envelope_json"])

            writer.writerow([
                row.get("id") or row.get("payload_id", ""),
                _format_timestamp(row.get("collected_at") or row.get("payload_collected_at") or row.get("received_at")),
                row.get("agent_id", ""),
                row.get("hostname", ""),
                row.get("username", ""),
                row.get("collector", ""),
                row.get("status") or row.get("validation_status", ""),
                payload_str,
            ])

        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


def _stream_json_logs(storage, filters: dict[str, Any], limit: int) -> Generator[str, None, None]:
    """Generator yielding JSON lines (NDJSON) for telemetry logs."""
    offset = 0
    exported = 0

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_logs(limit=batch_limit, offset=offset, **filters)
        if not rows:
            break

        for row in rows:
            # Symmetrically guarantee canonical aliases for consumers
            if "id" not in row and "payload_id" in row:
                row["id"] = row["payload_id"]
            if "collected_at" not in row and "payload_collected_at" in row:
                row["collected_at"] = row["payload_collected_at"]
            if "status" not in row and "validation_status" in row:
                row["status"] = row["validation_status"]
            yield json.dumps(row, default=str) + "\n"

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


def _stream_csv_threats(storage, limit: int) -> Generator[str, None, None]:
    """Generator yielding CSV rows for risk events."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([
        "id",
        "created_at",
        "payload_id",
        "agent_id",
        "username",
        "risk_level",
        "risk_score",
        "summary",
        "correlated_signals_json",
    ])
    yield buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)

    offset = 0
    exported = 0

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_risk_events(limit=batch_limit, offset=offset)
        if not rows:
            break

        for row in rows:
            signals_str = ""
            signals = row.get("correlated_signals_json")
            if isinstance(signals, (dict, list)):
                signals_str = json.dumps(signals)
            elif signals is not None:
                signals_str = str(signals)

            writer.writerow([
                row.get("id", ""),
                _format_timestamp(row.get("created_at")),
                row.get("payload_id", ""),
                row.get("agent_id", ""),
                row.get("username", ""),
                (row.get("risk_level") or "").upper(),
                row.get("risk_score", 0.0),
                row.get("summary", ""),
                signals_str,
            ])

        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


def _stream_json_threats(storage, limit: int) -> Generator[str, None, None]:
    """Generator yielding JSON lines for risk events."""
    offset = 0
    exported = 0

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_risk_events(limit=batch_limit, offset=offset)
        if not rows:
            break

        for row in rows:
            yield json.dumps(row, default=str) + "\n"

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


@router.get("/v1/export/logs")
@router.get("/export/logs")
@router.get("/v1/export/logs.csv")
@router.get("/export/logs.csv")
@router.get("/v1/export/logs.json")
@router.get("/export/logs.json")
async def export_logs(
    request: Request,
    format: str = Query("csv"),
    limit: int = Query(50000),
    agent_id: str | None = None,
    hostname: str | None = None,
    username: str | None = None,
    collector: str | None = None,
    status: str | None = None,
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Stream telemetry logs in CSV or JSON format with constant memory usage."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    path = request.url.path.lower()
    if path.endswith(".json") or path.endswith(".jsonl"):
        export_format = "json"
    elif path.endswith(".csv"):
        export_format = "csv"
    else:
        export_format = format.lower()

    bounded_limit = min(max(1, limit), 200000)

    filters = {
        k: v for k, v in {
            "agent_id": agent_id,
            "hostname": hostname,
            "username": username,
            "collector": collector,
            "status": status,
        }.items() if v
    }

    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    if export_format in ("json", "jsonl"):
        generator = _stream_json_logs(storage, filters, bounded_limit)
        mimetype = "application/x-ndjson"
        filename = f"insiedr_telemetry_{timestamp_str}.jsonl"
    else:
        generator = _stream_csv_logs(storage, filters, bounded_limit)
        mimetype = "text/csv; charset=utf-8"
        filename = f"insiedr_telemetry_{timestamp_str}.csv"

    return StreamingResponse(
        generator,
        media_type=mimetype,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-cache",
        },
    )


@router.get("/v1/export/threats")
@router.get("/export/threats")
@router.get("/v1/export/threats.csv")
@router.get("/export/threats.csv")
@router.get("/v1/export/threats.json")
@router.get("/export/threats.json")
async def export_threats(
    request: Request,
    format: str = Query("csv"),
    limit: int = Query(50000),
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Stream risk events and threat detections in CSV or JSON format."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    path = request.url.path.lower()
    if path.endswith(".json") or path.endswith(".jsonl"):
        export_format = "json"
    elif path.endswith(".csv"):
        export_format = "csv"
    else:
        export_format = format.lower()

    bounded_limit = min(max(1, limit), 200000)
    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    if export_format in ("json", "jsonl"):
        generator = _stream_json_threats(storage, bounded_limit)
        mimetype = "application/x-ndjson"
        filename = f"insiedr_threats_{timestamp_str}.jsonl"
    else:
        generator = _stream_csv_threats(storage, bounded_limit)
        mimetype = "text/csv; charset=utf-8"
        filename = f"insiedr_threats_{timestamp_str}.csv"

    return StreamingResponse(
        generator,
        media_type=mimetype,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-cache",
        },
    )


def _resolve_decrypted_payload(row: dict[str, Any]) -> dict[str, Any] | None:
    """Resolve decrypted telemetry payload from raw row data.
    
    Tries stored plaintext payload first, then decrypts encrypted_envelope_json on the fly
    using server crypto plugins so raw_logs and parameters are never null in exported datasets.
    """
    raw_payload_data = row.get("payload")
    if not raw_payload_data and row.get("payload_json"):
        try:
            raw_payload_data = json.loads(row["payload_json"])
        except Exception:
            raw_payload_data = None

    if isinstance(raw_payload_data, dict):
        return raw_payload_data

    # Attempt decrypting encrypted_envelope_json on the fly
    envelope = row.get("encrypted_envelope_json")
    if envelope:
        if isinstance(envelope, str):
            try:
                envelope = json.loads(envelope)
            except Exception:
                envelope = None
        if isinstance(envelope, dict):
            scheme = envelope.get("scheme") or row.get("crypto_scheme") or "aes-256-gcm"
            try:
                from server.plugin_registry import registry
                if not registry.schemes():
                    registry.initialize()
                plugin = registry.get(scheme)
                if plugin and hasattr(plugin, "decrypt"):
                    plaintext_bytes = plugin.decrypt(envelope)
                    from shared.protocol import parse_json_bytes
                    return parse_json_bytes(plaintext_bytes)
            except Exception:
                pass

    return None


def _extract_parameters(storage, row: dict[str, Any]) -> dict[str, Any]:
    """Extract model training feature parameters from a telemetry log record."""
    # 1. First extract on-the-fly from decrypted collector payload (Zero DB queries!)
    payload_obj = _resolve_decrypted_payload(row)

    if isinstance(payload_obj, dict):
        features: dict[str, Any] = {}
        for collector in payload_obj.get("collectors", []):
            if not isinstance(collector, dict) or collector.get("status") != "success":
                continue
            col_payload = collector.get("payload")
            if not isinstance(col_payload, dict):
                continue
            for key, value in col_payload.items():
                if isinstance(value, dict):
                    for nested_key, nested_value in value.items():
                        if isinstance(nested_value, dict):
                            for feature_name, feature_value in nested_value.items():
                                features.setdefault(feature_name, feature_value)
                        else:
                            features.setdefault(nested_key, nested_value)
                else:
                    features.setdefault(key, value)
        if features:
            return features

    # 2. Check if row already has pre-attached features
    if isinstance(row.get("features"), dict) and row["features"]:
        return row["features"]

    # 3. Last resort fallback to storage.get_feature_vector
    payload_id = row.get("id") or row.get("payload_id")
    if payload_id and hasattr(storage, "get_feature_vector") and callable(storage.get_feature_vector):
        try:
            stored_features = storage.get_feature_vector(payload_id)
            if stored_features and isinstance(stored_features, dict):
                return stored_features
        except Exception:
            pass

    return {}


def _build_excel_training_dataset(
    storage,
    filters: dict[str, Any],
    limit: int,
    days: int,
    username: str | None,
) -> bytes:
    """Build multi-tab Excel spreadsheet (.xlsx) for ML training dataset."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # Sheet 1: Dataset Overview (User, Collected At, Parameters, Raw Logs, Metadata)
    ws1 = wb.active
    ws1.title = "Dataset_Overview"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers_ws1 = [
        "User",
        "Collected At",
        "Parameters (Model Features)",
        "Raw Logs",
        "Agent ID",
        "Hostname",
        "Payload ID",
    ]
    ws1.append(headers_ws1)

    for col_idx in range(1, len(headers_ws1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws1.row_dimensions[1].height = 26
    ws1.freeze_panes = "A2"

    offset = 0
    exported = 0
    all_rows_data = []
    all_feature_keys = set()

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_logs(limit=batch_limit, offset=offset, **filters)
        if not rows:
            break

        for row in rows:
            user = row.get("username") or "unknown"
            collected_at = _format_timestamp(row.get("collected_at") or row.get("payload_collected_at") or row.get("received_at"))
            agent_id = row.get("agent_id", "")
            hostname = row.get("hostname", "")
            payload_id = row.get("id") or row.get("payload_id", "")

            # Raw logs
            raw_payload_data = _resolve_decrypted_payload(row)
            if raw_payload_data is None:
                raw_payload_data = row.get("encrypted_envelope_json")

            raw_logs_str = json.dumps(raw_payload_data, default=str) if isinstance(raw_payload_data, (dict, list)) else str(raw_payload_data or "")

            # Extracted model parameters
            parameters = _extract_parameters(storage, row)
            all_feature_keys.update(parameters.keys())
            parameters_str = json.dumps(parameters, default=str)

            ws1.append([
                user,
                collected_at,
                parameters_str,
                raw_logs_str,
                agent_id,
                hostname,
                payload_id,
            ])
            all_rows_data.append((user, collected_at, agent_id, hostname, parameters))

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break

    # Format data rows in Sheet 1
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    for row_idx in range(2, ws1.max_row + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers_ws1) + 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = left_align
            if is_even:
                cell.fill = zebra_fill

    ws1.column_dimensions["A"].width = 18
    ws1.column_dimensions["B"].width = 24
    ws1.column_dimensions["C"].width = 45
    ws1.column_dimensions["D"].width = 50
    ws1.column_dimensions["E"].width = 18
    ws1.column_dimensions["F"].width = 20
    ws1.column_dimensions["G"].width = 36

    # Sheet 2: Tabular Feature Matrix (Exploded parameter columns for ML training)
    ws2 = wb.create_sheet(title="Tabular_Feature_Matrix")
    sorted_features = sorted(list(all_feature_keys))
    headers_ws2 = ["User", "Collected At", "Agent ID", "Hostname"] + sorted_features
    ws2.append(headers_ws2)

    header2_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for col_idx in range(1, len(headers_ws2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header2_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws2.row_dimensions[1].height = 26
    ws2.freeze_panes = "E2"

    for user, collected_at, agent_id, hostname, params in all_rows_data:
        row_vals = [user, collected_at, agent_id, hostname]
        for f_name in sorted_features:
            val = params.get(f_name, 0.0)
            try:
                row_vals.append(float(val) if val is not None else 0.0)
            except (ValueError, TypeError):
                row_vals.append(str(val))
        ws2.append(row_vals)

    for row_idx in range(2, ws2.max_row + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(headers_ws2) + 1):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill

    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 24
    ws2.column_dimensions["C"].width = 18
    ws2.column_dimensions["D"].width = 20
    for col_idx, f_name in enumerate(sorted_features, start=5):
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = max(len(f_name) + 3, 14)

    # Sheet 3: Dataset Metadata
    ws3 = wb.create_sheet(title="Dataset_Metadata")
    meta_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws3.append(["Metadata Field", "Value", "Description"])
    for col_idx in range(1, 4):
        c = ws3.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = meta_fill
        c.alignment = center_align
    ws3.row_dimensions[1].height = 24

    metadata_entries = [
        ("Dataset Purpose", "InsiEDR Machine Learning Model Training (Domain IF, XGBoost, RedRVFL)", "Telemetry parameter and raw log dataset"),
        ("Export Timestamp (UTC)", datetime.now(timezone.utc).isoformat(), "Time when dataset was extracted"),
        ("Target User Filter", username or "All Users", "Filtered user or entire fleet"),
        ("Time Window (Days)", f"Last {days} Days", "Rolling temporal window for sample collection"),
        ("Total Telemetry Samples", len(all_rows_data), "Number of telemetry records in dataset"),
        ("Total Unique Feature Parameters", len(sorted_features), "Number of distinct model features in feature matrix"),
        ("InsiEDR Protocol Version", "2.0", "Sensor wire format version"),
        ("Usage in Python", "df = pd.read_excel('filename.xlsx', sheet_name='Tabular_Feature_Matrix')", "Ready for pandas / scikit-learn"),
    ]
    for m_row in metadata_entries:
        ws3.append(list(m_row))
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 40
    ws3.column_dimensions["C"].width = 55

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stream_csv_training_dataset(
    storage,
    filters: dict[str, Any],
    limit: int,
) -> Generator[str, None, None]:
    """Generator yielding CSV rows for ML model training dataset (User, Parameter, Raw Logs)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    # Header
    writer.writerow([
        "user",
        "collected_at",
        "parameters",
        "raw_logs",
        "agent_id",
        "hostname",
        "payload_id",
    ])
    yield buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)

    offset = 0
    exported = 0

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_logs(limit=batch_limit, offset=offset, **filters)
        if not rows:
            break

        for row in rows:
            user = row.get("username") or "unknown"
            collected_at = _format_timestamp(row.get("collected_at") or row.get("payload_collected_at") or row.get("received_at"))
            agent_id = row.get("agent_id", "")
            hostname = row.get("hostname", "")
            payload_id = row.get("id") or row.get("payload_id", "")

            # Raw logs
            raw_payload_data = _resolve_decrypted_payload(row)
            if raw_payload_data is None:
                raw_payload_data = row.get("encrypted_envelope_json")

            raw_logs_str = json.dumps(raw_payload_data, default=str) if isinstance(raw_payload_data, (dict, list)) else str(raw_payload_data or "")

            # Extracted model parameters
            parameters = _extract_parameters(storage, row)
            parameters_str = json.dumps(parameters, default=str)

            writer.writerow([
                user,
                collected_at,
                parameters_str,
                raw_logs_str,
                agent_id,
                hostname,
                payload_id,
            ])

        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


def _stream_json_training_dataset(
    storage,
    filters: dict[str, Any],
    limit: int,
) -> Generator[str, None, None]:
    """Generator yielding NDJSON lines for ML training dataset."""
    offset = 0
    exported = 0

    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_logs(limit=batch_limit, offset=offset, **filters)
        if not rows:
            break

        for row in rows:
            user = row.get("username") or "unknown"
            collected_at = _format_timestamp(row.get("collected_at") or row.get("payload_collected_at") or row.get("received_at"))
            agent_id = row.get("agent_id", "")
            hostname = row.get("hostname", "")
            payload_id = row.get("id") or row.get("payload_id", "")

            raw_payload_data = _resolve_decrypted_payload(row)
            if raw_payload_data is None:
                raw_payload_data = row.get("encrypted_envelope_json")

            parameters = _extract_parameters(storage, row)

            record = {
                "user": user,
                "collected_at": collected_at,
                "parameters": parameters,
                "raw_logs": raw_payload_data,
                "agent_id": agent_id,
                "hostname": hostname,
                "payload_id": payload_id,
            }
            yield json.dumps(record, default=str) + "\n"

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


@router.get("/v1/export/training-dataset")
@router.get("/export/training-dataset")
@router.get("/v1/export/training-dataset.xlsx")
@router.get("/export/training-dataset.xlsx")
@router.get("/v1/export/training-dataset.csv")
@router.get("/export/training-dataset.csv")
@router.get("/v1/export/training-dataset.json")
@router.get("/export/training-dataset.json")
async def export_training_dataset(
    request: Request,
    username: str | None = Query(None, description="Target user (e.g. 'alice') or omit for all users"),
    days: int = Query(30, ge=1, le=365, description="Number of historical days (N) to include in dataset"),
    format: str = Query("xlsx", description="Format: 'xlsx' (Excel), 'csv', or 'json'"),
    limit: int = Query(100000, ge=1, description="Maximum number of log samples to export (default 100000, or higher for complete dataset)"),
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Export training dataset containing User, Parameters, and Raw Logs over N days.
    
    Specifically engineered for training InsiEDR machine learning models (Domain Isolation Forest,
    Scenario XGBoost, and RedRVFL sequence model).
    """
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    path = request.url.path.lower()
    if path.endswith(".xlsx"):
        export_format = "xlsx"
    elif path.endswith(".csv"):
        export_format = "csv"
    elif path.endswith(".json") or path.endswith(".jsonl"):
        export_format = "json"
    else:
        export_format = format.lower()

    if export_format == "xlsx":
        bounded_limit = min(limit, 50000)
    else:
        bounded_limit = max(1, limit)
    start_time = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()

    filters: dict[str, Any] = {"start_time": start_time}
    if username and username.strip():
        filters["username"] = username.strip()

    user_label = username.strip() if username and username.strip() else "all_users"
    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    if export_format == "xlsx":
        try:
            excel_bytes = await anyio.to_thread.run_sync(
                lambda: _build_excel_training_dataset(
                    storage=storage,
                    filters=filters,
                    limit=bounded_limit,
                    days=days,
                    username=username,
                )
            )
            filename = f"insiedr_training_dataset_{user_label}_{days}d_{timestamp_str}.xlsx"
            return Response(
                content=excel_bytes,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f'attachment; filename="{filename}"',
                    "Cache-Control": "no-cache",
                },
            )
        except Exception as exc:
            import logging
            logging.getLogger("insiedr.export").error("Failed to generate Excel dataset: %s", exc, exc_info=True)
            return JSONResponse({"ok": False, "error": f"Failed to generate Excel dataset: {exc}"}, status_code=500)

    elif export_format in ("json", "jsonl"):
        generator = _stream_json_training_dataset(storage, filters, bounded_limit)
        filename = f"insiedr_training_dataset_{user_label}_{days}d_{timestamp_str}.jsonl"
        return StreamingResponse(
            generator,
            media_type="application/x-ndjson",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-cache",
            },
        )
    else:
        generator = _stream_csv_training_dataset(storage, filters, bounded_limit)
        filename = f"insiedr_training_dataset_{user_label}_{days}d_{timestamp_str}.csv"
        return StreamingResponse(
            generator,
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-cache",
            },
        )


# --------------------------------------------------------------------------
# Normalized Features & Keystrokes Export (Excel .xlsx & CSV)
# --------------------------------------------------------------------------


def _stream_csv_features(storage, filters: dict[str, Any], limit: int) -> Generator[str, None, None]:
    """Generator yielding CSV rows for normalized features."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([
        "id",
        "payload_id",
        "collected_at",
        "agent_id",
        "username",
        "hostname",
        "collector",
        "feature_name",
        "value",
        "source_quality",
    ])
    yield buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)

    offset = 0
    exported = 0
    while exported < limit:
        batch_limit = min(CHUNK_BATCH_SIZE, limit - exported)
        rows = storage.list_normalized_features(limit=batch_limit, offset=offset, **filters) if hasattr(storage, "list_normalized_features") else []
        if not rows:
            break

        for row in rows:
            val = row.get("feature_value_numeric")
            if val is None:
                val = row.get("feature_value_text")
            if val is None and row.get("feature_value_json") is not None:
                val = json.dumps(row["feature_value_json"])

            writer.writerow([
                row.get("id", ""),
                row.get("payload_id", ""),
                _format_timestamp(row.get("feature_timestamp") or row.get("created_at")),
                row.get("agent_id", ""),
                row.get("username", ""),
                row.get("hostname", ""),
                row.get("collector", ""),
                row.get("feature_name", ""),
                str(val if val is not None else ""),
                row.get("source_quality", ""),
            ])

        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)

        count = len(rows)
        exported += count
        offset += count
        if count < batch_limit:
            break


def _build_excel_features(storage, filters: dict[str, Any], limit: int) -> bytes:
    """Build styled Excel workbook (.xlsx) for normalized features."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Normalized_Features"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = [
        "Feature ID",
        "Payload ID",
        "Timestamp",
        "Agent ID",
        "Username",
        "Hostname",
        "Collector",
        "Feature Name",
        "Numeric Value",
        "Text / JSON Value",
        "Source Quality",
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    rows = storage.list_normalized_features(limit=limit, offset=0, **filters) if hasattr(storage, "list_normalized_features") else []
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, r in enumerate(rows, start=2):
        num_val = r.get("feature_value_numeric")
        txt_val = r.get("feature_value_text")
        if txt_val is None and r.get("feature_value_json") is not None:
            txt_val = json.dumps(r["feature_value_json"])

        ws.append([
            r.get("id", ""),
            r.get("payload_id", ""),
            _format_timestamp(r.get("feature_timestamp") or r.get("created_at")),
            r.get("agent_id", ""),
            r.get("username", ""),
            r.get("hostname", ""),
            r.get("collector", ""),
            r.get("feature_name", ""),
            num_val if num_val is not None else "",
            str(txt_val or ""),
            r.get("source_quality", ""),
        ])
        is_even = (idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=idx, column=col_idx)
            c.border = thin_border
            c.alignment = left_align
            if is_even:
                c.fill = zebra_fill

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col[:100])
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_excel_keystrokes(storage, limit: int, username: str | None = None, agent_id: str | None = None) -> bytes:
    """Build Excel workbook (.xlsx) dedicated to Keystroke Dynamics biometrics."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Keystroke_Dynamics"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    headers = [
        "Username",
        "Hostname",
        "Agent ID",
        "Collected At",
        "Mean Flight (ms)",
        "Std Flight (ms)",
        "Mean Dwell (ms)",
        "Std Dwell (ms)",
        "Typing Speed (CPM)",
        "Backspace Ratio",
        "Raw Keystroke Timings [dwell, flight]",
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "E2"

    collector_rows = storage.list_collector_results(limit=limit, offset=0, collector="keystroke-collector", username=username) if hasattr(storage, "list_collector_results") else []
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for idx, r in enumerate(collector_rows, start=2):
        p = r.get("payload_json") or {}
        if isinstance(p, str):
            try:
                p = json.loads(p)
            except Exception:
                p = {}
        features = r.get("features_json") or {}
        if isinstance(features, str):
            try:
                features = json.loads(features)
            except Exception:
                features = {}

        mean_flight = p.get("mean_flight_time_ms") or features.get("mean_flight_time_ms") or ""
        std_flight = p.get("std_flight_time_ms") or features.get("std_flight_time_ms") or ""
        mean_dwell = p.get("mean_dwell_time_ms") or features.get("mean_dwell_time_ms") or ""
        std_dwell = p.get("std_dwell_time_ms") or features.get("std_dwell_time_ms") or ""
        speed_cpm = p.get("typing_speed_cpm") or features.get("typing_speed_cpm") or ""
        backspace = p.get("backspace_ratio") or features.get("backspace_ratio") or ""
        raw_timings = json.dumps(p.get("keystroke_timings") or [])

        ws.append([
            r.get("username", ""),
            r.get("hostname", ""),
            r.get("agent_id", ""),
            _format_timestamp(r.get("collector_collected_at") or r.get("received_at")),
            mean_flight,
            std_flight,
            mean_dwell,
            std_dwell,
            speed_cpm,
            backspace,
            raw_timings,
        ])
        is_even = (idx % 2 == 0)
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=idx, column=col_idx)
            c.border = thin_border
            c.alignment = left_align
            if is_even:
                c.fill = zebra_fill

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col[:100])
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _stream_csv_keystrokes(storage, limit: int, username: str | None = None) -> Generator[str, None, None]:
    """Generator yielding CSV rows dedicated to Keystroke Dynamics."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([
        "username",
        "hostname",
        "agent_id",
        "collected_at",
        "mean_flight_time_ms",
        "std_flight_time_ms",
        "mean_dwell_time_ms",
        "std_dwell_time_ms",
        "typing_speed_cpm",
        "backspace_ratio",
        "keystroke_timings",
    ])
    yield buffer.getvalue()
    buffer.seek(0)
    buffer.truncate(0)

    collector_rows = storage.list_collector_results(limit=limit, offset=0, collector="keystroke-collector", username=username) if hasattr(storage, "list_collector_results") else []
    for r in collector_rows:
        p = r.get("payload_json") or {}
        if isinstance(p, str):
            try:
                p = json.loads(p)
            except Exception:
                p = {}
        features = r.get("features_json") or {}
        if isinstance(features, str):
            try:
                features = json.loads(features)
            except Exception:
                features = {}

        writer.writerow([
            r.get("username", ""),
            r.get("hostname", ""),
            r.get("agent_id", ""),
            _format_timestamp(r.get("collector_collected_at") or r.get("received_at")),
            p.get("mean_flight_time_ms") or features.get("mean_flight_time_ms") or "",
            p.get("std_flight_time_ms") or features.get("std_flight_time_ms") or "",
            p.get("mean_dwell_time_ms") or features.get("mean_dwell_time_ms") or "",
            p.get("std_dwell_time_ms") or features.get("std_dwell_time_ms") or "",
            p.get("typing_speed_cpm") or features.get("typing_speed_cpm") or "",
            p.get("backspace_ratio") or features.get("backspace_ratio") or "",
            json.dumps(p.get("keystroke_timings") or []),
        ])
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)


@router.get("/v1/export/features")
@router.get("/export/features")
@router.get("/v1/export/features.xlsx")
@router.get("/export/features.xlsx")
@router.get("/v1/export/features.csv")
@router.get("/export/features.csv")
async def export_features(
    request: Request,
    format: str = Query("csv"),
    limit: int = Query(50000),
    agent_id: str | None = None,
    username: str | None = None,
    collector: str | None = None,
    feature_name: str | None = None,
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Export normalized features from all 20 sensors in Excel (.xlsx) or CSV format."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    path = request.url.path.lower()
    if path.endswith(".xlsx"):
        export_format = "xlsx"
    elif path.endswith(".csv"):
        export_format = "csv"
    else:
        export_format = format.lower()

    bounded_limit = min(max(1, limit), 200000)
    filters = {
        k: v for k, v in {
            "agent_id": agent_id,
            "username": username,
            "collector": collector,
            "feature_name": feature_name,
        }.items() if v
    }

    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    if export_format == "xlsx":
        try:
            excel_bytes = await anyio.to_thread.run_sync(
                lambda: _build_excel_features(storage, filters, bounded_limit)
            )
            filename = f"insiedr_features_{timestamp_str}.xlsx"
            return Response(
                content=excel_bytes,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f'attachment; filename="{filename}"',
                    "Cache-Control": "no-cache",
                },
            )
        except Exception as exc:
            return JSONResponse({"ok": False, "error": f"Failed to generate Excel features: {exc}"}, status_code=500)
    else:
        generator = _stream_csv_features(storage, filters, bounded_limit)
        filename = f"insiedr_features_{timestamp_str}.csv"
        return StreamingResponse(
            generator,
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-cache",
            },
        )


@router.get("/v1/export/keystrokes")
@router.get("/export/keystrokes")
@router.get("/v1/export/keystrokes.xlsx")
@router.get("/export/keystrokes.xlsx")
@router.get("/v1/export/keystrokes.csv")
@router.get("/export/keystrokes.csv")
async def export_keystrokes(
    request: Request,
    format: str = Query("xlsx"),
    limit: int = Query(50000),
    username: str | None = None,
    agent_id: str | None = None,
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Export keystroke dynamics biometric feature logs in Excel (.xlsx) or CSV format."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    path = request.url.path.lower()
    if path.endswith(".xlsx"):
        export_format = "xlsx"
    elif path.endswith(".csv"):
        export_format = "csv"
    else:
        export_format = format.lower()

    bounded_limit = min(max(1, limit), 200000)
    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    if export_format == "xlsx":
        try:
            excel_bytes = await anyio.to_thread.run_sync(
                lambda: _build_excel_keystrokes(storage, bounded_limit, username=username, agent_id=agent_id)
            )
            filename = f"insiedr_keystroke_biometrics_{timestamp_str}.xlsx"
            return Response(
                content=excel_bytes,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f'attachment; filename="{filename}"',
                    "Cache-Control": "no-cache",
                },
            )
        except Exception as exc:
            return JSONResponse({"ok": False, "error": f"Failed to generate Excel keystrokes: {exc}"}, status_code=500)
    else:
        generator = _stream_csv_keystrokes(storage, bounded_limit, username=username)
        filename = f"insiedr_keystroke_biometrics_{timestamp_str}.csv"
        return StreamingResponse(
            generator,
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-cache",
            },
        )


# ------------------------------------------------------------------------------
# Collector-Specific Dataset Export & Dynamic Feature Schema
# ------------------------------------------------------------------------------

def flatten_collector_payload(payload: Any, prefix: str = "", sep: str = "_") -> dict[str, Any]:
    """Recursively flatten nested dictionary features into single-level column keys.
    Preserves types: int, float, bool, str.
    Example:
        {"event_count": 15, "etw_lifecycle": {"started_count": 5}}
        -> {"event_count": 15, "etw_lifecycle_started_count": 5}
    """
    flat: dict[str, Any] = {}
    if not isinstance(payload, dict):
        return flat
    for k, v in payload.items():
        if str(k).startswith("_"):
            continue
        clean_key = str(k).strip().replace(" ", "_").replace("-", "_").replace(".", "_")
        col_name = f"{prefix}{sep}{clean_key}" if prefix else clean_key
        if isinstance(v, dict):
            flat.update(flatten_collector_payload(v, prefix=col_name, sep=sep))
        elif isinstance(v, (list, tuple)):
            if not v:
                flat[col_name] = ""
            elif all(isinstance(x, (int, float, str, bool)) for x in v):
                flat[col_name] = ";".join(str(x) for x in v)
            else:
                for index, item in enumerate(v):
                    flat.update(flatten_collector_payload({str(index): item}, prefix=col_name, sep=sep))
        elif isinstance(v, bool):
            flat[col_name] = v
        elif isinstance(v, (int, float)):
            flat[col_name] = v
        elif v is None:
            flat[col_name] = ""
        else:
            flat[col_name] = str(v)
    return flat


def _normalize_iso_date_bound(val: str | None, is_end: bool = False) -> str | None:
    if not val or not val.strip():
        return None
    s = val.strip()
    if len(s) == 10 and s.count("-") == 2:
        return f"{s}T23:59:59Z" if is_end else f"{s}T00:00:00Z"
    return s


def _build_collector_filename(collector: str, start_date: str | None, end_date: str | None, username: str | None) -> str:
    clean_col = re.sub(r"[^\w\-]", "_", (collector or "collector").strip().lower())
    if username and username.strip():
        sanitized_user = re.sub(r"[^\w\-]", "_", username.strip())
        clean_u = f"_{sanitized_user}"
    else:
        clean_u = ""
    clean_start = start_date.strip().split("T")[0] if start_date and start_date.strip() else None
    clean_end = end_date.strip().split("T")[0] if end_date and end_date.strip() else None

    if clean_start and clean_end:
        return f"{clean_col}{clean_u}_{clean_start}_to_{clean_end}.csv"
    elif clean_start:
        return f"{clean_col}{clean_u}_from_{clean_start}.csv"
    elif clean_end:
        return f"{clean_col}{clean_u}_until_{clean_end}.csv"
    else:
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        return f"{clean_col}{clean_u}_dataset_{ts}.csv"


def _stream_csv_collector_dataset(
    storage,
    collector: str,
    filters: dict[str, Any],
    limit: int,
) -> Generator[str, None, None]:
    """Stream CSV rows for a specific collector with dynamic column unrolling."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    base_columns = ["timestamp", "username", "hostname", "agent_id", "status"]
    call_filters = dict(filters)
    call_filters["collector"] = collector

    # Discover all columns while spooling rows to disk after 1 MiB. This avoids
    # silently dropping fields introduced after the first batch and keeps RAM
    # bounded independently of the number of exported records.
    from tempfile import SpooledTemporaryFile
    with SpooledTemporaryFile(max_size=1024 * 1024, mode="w+t", encoding="utf-8") as spool:
        columns = set()
        offset = 0
        while offset < limit:
            batch_limit = min(CHUNK_BATCH_SIZE, limit - offset)
            rows = storage.list_collector_results(limit=batch_limit, offset=offset, **call_filters)
            if not rows:
                break
            for record in rows:
                payload = record.get("payload") or record.get("payload_json") or {}
                if isinstance(payload, str):
                    try:
                        payload = json.loads(payload)
                    except ValueError:
                        payload = {}
                flat = flatten_collector_payload(payload)
                columns.update(flat)
                if len(columns) > 10000:
                    raise ValueError("Collector export exceeds 10000 distinct feature columns")
                base = [_format_timestamp(record.get("collected_at") or record.get("collector_collected_at") or record.get("received_at")),
                        record.get("username") or "", record.get("hostname") or "",
                        record.get("agent_id") or "", record.get("status") or "success"]
                spool.write(json.dumps([base, flat], default=str) + "\n")
            offset += len(rows)
            if len(rows) < batch_limit:
                break
        feature_columns = sorted(columns)
        writer.writerow(base_columns + feature_columns)
        yield buffer.getvalue()
        buffer.seek(0)
        buffer.truncate(0)
        spool.seek(0)
        for index, line in enumerate(spool):
            base, flat = json.loads(line)
            writer.writerow(base + [flat.get(column, "") for column in feature_columns])
            if (index + 1) % 100 == 0:
                yield buffer.getvalue()
                buffer.seek(0)
                buffer.truncate(0)
        if buffer.tell():
            yield buffer.getvalue()


@router.get("/v1/export/collectors")
@router.get("/export/collectors")
def get_export_collectors(
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Return available telemetry collectors and known users for export filtering."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    collectors = getattr(storage, "get_distinct_collectors", lambda: [
        "logon", "file", "process", "network", "device", "http",
        "lsass_monitor", "registry", "keystroke-collector"
    ])()
    usernames = getattr(storage, "get_distinct_usernames", lambda: [])()

    return {
        "ok": True,
        "collectors": collectors,
        "usernames": usernames,
    }


@router.get("/v1/export/collector-preview")
@router.get("/export/collector-preview")
def preview_collector_dataset(
    collector: str = Query(..., description="Target collector name (e.g. 'logon', 'file', 'process')"),
    start_date: str | None = Query(None, description="Start date (YYYY-MM-DD or ISO timestamp)"),
    end_date: str | None = Query(None, description="End date (YYYY-MM-DD or ISO timestamp)"),
    username: str | None = Query(None, description="Filter by user or omit for all"),
    limit: int = Query(50, ge=1, le=200, description="Preview sample limit"),
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Preview collector-specific dataset schema and matching records before full export."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    if not collector or not collector.strip():
        return JSONResponse({"ok": False, "error": "Collector parameter is required"}, status_code=400)

    start_iso = _normalize_iso_date_bound(start_date, is_end=False)
    end_iso = _normalize_iso_date_bound(end_date, is_end=True)

    filters: dict[str, Any] = {"collector": collector.strip()}
    if username and username.strip():
        filters["username"] = username.strip()
    if start_iso:
        filters["start_time"] = start_iso
    if end_iso:
        filters["end_time"] = end_iso

    rows = storage.list_collector_results(limit=limit, offset=0, **filters)

    base_columns = ["timestamp", "username", "hostname", "agent_id", "status"]
    discovered_keys = set()
    flattened_rows = []

    for r in rows:
        raw_payload = r.get("payload") or r.get("payload_json") or {}
        if isinstance(raw_payload, str):
            try:
                raw_payload = json.loads(raw_payload)
            except Exception:
                raw_payload = {}
        flat = flatten_collector_payload(raw_payload)
        discovered_keys.update(flat.keys())
        ts = _format_timestamp(r.get("collected_at") or r.get("collector_collected_at") or r.get("received_at"))

        item = {
            "timestamp": ts,
            "username": r.get("username") or "",
            "hostname": r.get("hostname") or "",
            "agent_id": r.get("agent_id") or "",
            "status": r.get("status") or "success",
        }
        item.update(flat)
        flattened_rows.append(item)

    feature_columns = sorted(list(discovered_keys))
    all_columns = base_columns + feature_columns

    # Normalize rows to ensure every column exists
    preview_rows = []
    for r in flattened_rows:
        row_dict = {}
        for c in all_columns:
            row_dict[c] = r.get(c, "")
        preview_rows.append(row_dict)

    return {
        "ok": True,
        "collector": collector.strip(),
        "total_samples": len(rows),
        "columns": all_columns,
        "feature_columns": feature_columns,
        "preview_rows": preview_rows,
        "filename": _build_collector_filename(collector, start_date, end_date, username),
    }


@router.get("/v1/export/collector-dataset.csv")
@router.get("/v1/export/collector-dataset")
@router.get("/export/collector-dataset.csv")
@router.get("/export/collector-dataset")
async def export_collector_dataset(
    request: Request,
    collector: str = Query(..., description="Target collector name (e.g. 'logon', 'file', 'process')"),
    start_date: str | None = Query(None, description="Start date (YYYY-MM-DD or ISO timestamp)"),
    end_date: str | None = Query(None, description="End date (YYYY-MM-DD or ISO timestamp)"),
    username: str | None = Query(None, description="Filter by user or omit for all"),
    limit: int = Query(50000, ge=1, le=100000, description="Maximum number of log samples to export"),
    storage=Depends(get_storage),
    operator: OperatorPrincipal = Depends(require_operator("operator:read")),
):
    """Export collector-specific dataset with unrolled feature columns in CSV format."""
    if storage is None:
        return JSONResponse({"ok": False, "error": "Storage is not configured"}, status_code=503)

    if not collector or not collector.strip():
        return JSONResponse({"ok": False, "error": "Collector parameter is required"}, status_code=400)

    bounded_limit = min(max(1, limit), 200000)
    start_iso = _normalize_iso_date_bound(start_date, is_end=False)
    end_iso = _normalize_iso_date_bound(end_date, is_end=True)

    filters: dict[str, Any] = {}
    if username and username.strip():
        filters["username"] = username.strip()
    if start_iso:
        filters["start_time"] = start_iso
    if end_iso:
        filters["end_time"] = end_iso

    filename = _build_collector_filename(collector, start_date, end_date, username)
    generator = _stream_csv_collector_dataset(storage, collector.strip(), filters, bounded_limit)

    return StreamingResponse(
        generator,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-cache",
        },
    )
